from __future__ import annotations

from pathlib import Path
from textwrap import wrap

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as RLImage,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(r"d:\test_workspace")
DOC_OUT = ROOT / "output" / "documents"
PDF_OUT = ROOT / "output" / "pdf"
MEDIA_DIR = DOC_OUT / "original_media"
RENDER_DIR = PDF_OUT / "rendered_attachment_pricing"

FRONT_FLOWCHART_PATH = DOC_OUT / "附加项新定价_前置配置流程图.png"
FLOWCHART_PATH = DOC_OUT / "附加项新定价_流程图.png"
DOCX_PATH = DOC_OUT / "附加项新定价_优化版.docx"
PDF_PATH = PDF_OUT / "附加项新定价_优化版.pdf"

FONT_REGULAR = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")


TITLE = "附加项新定价配置说明"
SUBTITLE = "按上架国家与会员等级维护附加项服务定价"
DOC_DATE = "2026-05-26"


FRONT_STEP_ROWS = [
    ("前置步骤", "配置内容", "输出结果"),
    (
        "1. 新增/编辑会员等级",
        "进入 客户管理 > 会员等级，维护等级名称、等级标识、多语言、成员条件、等级权重、等级说明、状态与排序。会员等级页只维护等级基础属性。",
        "形成可被国家会员定价、附加项定价引用的会员等级基础数据。",
    ),
    (
        "2. 配置全球会员定价",
        "进入 客户管理 > 全球会员定价，按国家编辑会员等级定价体系；配置会员套餐周期、会员价格、划线价，以及代采阶梯手续费。",
        "同一会员等级可以在不同国家维护独立会员费和代采费体系。",
    ),
    (
        "3. 会员中心展示",
        "前台 Service/Rates 页面展示当前国家下已配置、已启用的会员等级卡片，包含等级名称、套餐周期、会员价格、划线价和代采服务费比例。",
        "客户可看到新配置会员等级；当前账号等级显示 Current Level 标识。",
    ),
]


SUMMARY = [
    "普通附加项的新增、编辑入口及基础资料字段保持原有逻辑不变。",
    "新增定价维护维度：上架国家、附加项类型、服务项、适用会员等级、计费方式、成本价、客户价、状态与排序。",
    "同一附加项可在不同国家、不同会员等级下维护独立客户价。",
    "前台按当前站点国家和登录账号会员等级过滤展示；未配置、未启用或会员等级不匹配的附加项不展示、不可选。",
]


RULE_ROWS = [
    ("配置维度", "规则说明", "前台影响"),
    ("上架国家", "普通附加项仍在基础资料中选择上架国家，进入对应国家后再维护定价。", "客户所在站点国家未覆盖时，不展示该附加项。"),
    ("会员等级", "在国家定价中选择适用等级，如一般会员、VIP会员、企业会员。", "仅当前会员等级匹配的客户可见、可选。"),
    ("价格配置", "按计费方式维护成本价与客户价，支持普通计费、起订量、按套、按时、面议、阶梯价、仓储等方式。", "前台展示该等级对应的客户价。"),
    ("状态配置", "应用状态控制是否可使用，记录状态控制该定价是否启用。", "禁用或不可用的记录不进入前台选择范围。"),
    ("通用价", "如系统允许适用等级留空，可作为通用价格；若要严格按等级区分，应分别配置各等级价格。", "避免通用价覆盖精细化会员价预期。"),
]


STEPS = [
    ("1. 维护普通附加项基础资料", "进入 B2B系统 > 附加项 > 普通附加项，新增或编辑附加项名称、分类、应用场景、计费单位等字段。此部分功能保持原有逻辑。"),
    ("2. 选择上架国家", "在附加项基础资料中选择上架国家。一个附加项可以覆盖多个国家，后续需要分别进入国家定价维护价格。"),
    ("3. 进入全球服务定价", "在 附加项 > 全球服务定价 中选择目标国家并点击编辑，进入该国家的质检定价和附加项定价页面。"),
    ("4. 新增或编辑附加项定价", "在附加项定价页选择商品附加项或发货附加项，点击新增定价，选择服务项、适用等级、计费方式、价格、状态和排序。"),
    ("5. 前台展示校验", "使用对应会员等级账号进入购物车附加项选择页，确认可见附加项和会员价正确；非对应等级账号不应看到未配置的附加项。"),
]


CHECKS = [
    "新增普通附加项后，基础资料保存、编辑和上架国家选择均可正常使用。",
    "每个启用国家都可以进入全球服务定价页面并维护附加项定价。",
    "同一国家、同一服务项可按不同会员等级维护不同客户价。",
    "前台 VIP 账号可以看到已配置 VIP 等级的附加项及对应价格。",
    "前台一般会员账号无法看到仅配置给 VIP 的附加项。",
    "禁用、不可用或记录状态未启用的定价不在前台展示。",
]


SCREENSHOTS = [
    ("图 1 普通附加项新增入口与基础资料", MEDIA_DIR / "image1.png"),
    ("图 2 选择附加项上架国家", MEDIA_DIR / "image2.png"),
    ("图 3 全球服务定价国家列表入口", MEDIA_DIR / "image3.png"),
    ("图 4 新增/编辑国家下的商品附加项服务定价", MEDIA_DIR / "image4.png"),
    ("图 5 选择适用会员等级", MEDIA_DIR / "image5.png"),
    ("图 6 按会员等级筛选后产生对应附加项定价记录", MEDIA_DIR / "image6.png"),
    ("图 7 VIP 会员前台可选择已配置附加项", MEDIA_DIR / "image8.png"),
    ("图 8 一般会员未配置该等级定价时不展示该附加项", MEDIA_DIR / "image9.png"),
]


def ensure_dirs() -> None:
    DOC_OUT.mkdir(parents=True, exist_ok=True)
    PDF_OUT.mkdir(parents=True, exist_ok=True)
    RENDER_DIR.mkdir(parents=True, exist_ok=True)


def get_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size=size)


def wrap_zh(text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for ch in text:
        trial = current + ch
        if font.getlength(trial) <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str,
    max_width: int,
    line_gap: int = 8,
) -> None:
    lines: list[str] = []
    for part in text.split("\n"):
        lines.extend(wrap_zh(part, font, max_width))
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = box[1] + (box[3] - box[1] - total_h) / 2
    for line, h in zip(lines, heights):
        line_w = font.getlength(line)
        x = box[0] + (box[2] - box[0] - line_w) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += h + line_gap


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str) -> None:
    draw.line([start, end], fill=color, width=5)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex >= sx else -1
        points = [(ex, ey), (ex - direction * 18, ey - 10), (ex - direction * 18, ey + 10)]
    else:
        direction = 1 if ey >= sy else -1
        points = [(ex, ey), (ex - 10, ey - direction * 18), (ex + 10, ey - direction * 18)]
    draw.polygon(points, fill=color)


def generate_front_flowchart() -> None:
    img = Image.new("RGB", (1800, 620), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(40, bold=True)
    box_font = get_font(28, bold=True)
    small_font = get_font(23)

    draw.rectangle((0, 0, 1800, 620), fill="#F8FAFC")
    draw.text((70, 42), "前置配置流程：会员等级与会员定价", font=title_font, fill="#0B2545")
    draw.text((70, 96), "先完成会员等级和各国会员定价，再进入附加项按等级维护服务价格。", font=small_font, fill="#475569")

    boxes = [
        (80, 190, 410, 330, "新增/编辑\n会员等级", "#EEF2FF", "#4F46E5"),
        (520, 190, 850, 330, "全球会员定价\n按国家维护会员费", "#ECFEFF", "#0891B2"),
        (960, 190, 1290, 330, "配置代采费\n阶梯手续费", "#FFF7ED", "#EA580C"),
        (1400, 190, 1730, 330, "会员中心展示\n新会员等级", "#ECFDF5", "#16A34A"),
    ]
    for x1, y1, x2, y2, text, fill, edge in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=24, fill=fill, outline=edge, width=4)
        draw_centered_text(draw, (x1 + 18, y1 + 10, x2 - 18, y2 - 10), text, box_font, "#0B2545", x2 - x1 - 36)

    for left, right in zip(boxes, boxes[1:]):
        arrow(draw, (left[2] + 18, 260), (right[0] - 18, 260), "#2563EB")

    note = (360, 430, 1440, 545)
    draw.rounded_rectangle(note, radius=20, fill="#FFFFFF", outline="#CBD5E1", width=3)
    draw_centered_text(
        draw,
        note,
        "关键关系：会员等级是基础维度；全球会员定价决定各国会员费与代采费；附加项定价再引用会员等级维护服务价格。",
        small_font,
        "#334155",
        1000,
    )
    img.save(FRONT_FLOWCHART_PATH, quality=95)


def generate_flowchart() -> None:
    img = Image.new("RGB", (1800, 780), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(42, bold=True)
    box_font = get_font(27, bold=True)
    small_font = get_font(24)
    label_font = get_font(24, bold=True)

    draw.rectangle((0, 0, 1800, 780), fill="#F7FAFC")
    draw.text((70, 42), "附加项新定价配置流程", font=title_font, fill="#0B2545")
    draw.text((70, 96), "后台按国家与会员等级维护价格，前台按当前账号匹配展示。", font=small_font, fill="#475569")

    box_fill = "#E8F1FF"
    box_edge = "#2563EB"
    boxes = [
        (70, 180, 330, 300, "普通附加项\n新增/编辑不变"),
        (390, 180, 650, 300, "选择\n上架国家"),
        (710, 180, 970, 300, "进入\n全球服务定价"),
        (1030, 180, 1290, 300, "新增/编辑\n服务定价"),
        (1350, 180, 1610, 300, "配置等级\n价格与状态"),
    ]
    for box in boxes:
        x1, y1, x2, y2, text = box
        draw.rounded_rectangle((x1, y1, x2, y2), radius=24, fill=box_fill, outline=box_edge, width=4)
        draw_centered_text(draw, (x1 + 16, y1 + 10, x2 - 16, y2 - 10), text, box_font, "#0B2545", x2 - x1 - 32)

    for left, right in zip(boxes, boxes[1:]):
        arrow(draw, (left[2] + 12, 240), (right[0] - 16, 240), "#2563EB")

    decision = [(900, 420), (1060, 330), (1220, 420), (1060, 510)]
    draw.polygon(decision, fill="#FFF7ED", outline="#F97316")
    draw.line(decision + [decision[0]], fill="#F97316", width=4)
    draw_centered_text(draw, (925, 355, 1195, 485), "当前会员等级\n有启用定价?", box_font, "#7C2D12", 250)
    arrow(draw, (1480, 304), (1120, 365), "#2563EB")

    show_box = (1300, 560, 1660, 690)
    hide_box = (480, 560, 840, 690)
    draw.rounded_rectangle(show_box, radius=24, fill="#ECFDF5", outline="#16A34A", width=4)
    draw.rounded_rectangle(hide_box, radius=24, fill="#FEF2F2", outline="#DC2626", width=4)
    draw_centered_text(draw, show_box, "展示附加项\n显示对应会员价", box_font, "#14532D", 320)
    draw_centered_text(draw, hide_box, "不展示 / 不可选\n避免跨等级使用", box_font, "#7F1D1D", 320)

    arrow(draw, (1185, 455), (1300, 620), "#16A34A")
    draw.text((1215, 510), "是", font=label_font, fill="#15803D")
    arrow(draw, (935, 455), (840, 620), "#DC2626")
    draw.text((855, 510), "否", font=label_font, fill="#B91C1C")

    img.save(FLOWCHART_PATH, quality=95)


def register_pdf_fonts() -> None:
    pdfmetrics.registerFont(TTFont("MSYH", str(FONT_REGULAR)))
    pdfmetrics.registerFont(TTFont("MSYH-Bold", str(FONT_BOLD)))


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "title",
            parent=base["Title"],
            fontName="MSYH-Bold",
            fontSize=26,
            leading=34,
            textColor=colors.HexColor("#0B2545"),
            alignment=TA_LEFT,
            spaceAfter=10,
        ),
        "subtitle": ParagraphStyle(
            "subtitle",
            parent=base["Normal"],
            fontName="MSYH",
            fontSize=12.5,
            leading=18,
            textColor=colors.HexColor("#475569"),
            spaceAfter=16,
            wordWrap="CJK",
        ),
        "h1": ParagraphStyle(
            "h1",
            parent=base["Heading1"],
            fontName="MSYH-Bold",
            fontSize=16,
            leading=22,
            textColor=colors.HexColor("#1D4ED8"),
            spaceBefore=8,
            spaceAfter=8,
        ),
        "h2": ParagraphStyle(
            "h2",
            parent=base["Heading2"],
            fontName="MSYH-Bold",
            fontSize=13,
            leading=18,
            textColor=colors.HexColor("#0F172A"),
            spaceBefore=8,
            spaceAfter=6,
        ),
        "body": ParagraphStyle(
            "body",
            parent=base["BodyText"],
            fontName="MSYH",
            fontSize=10.3,
            leading=15,
            textColor=colors.HexColor("#1F2937"),
            wordWrap="CJK",
            spaceAfter=5,
        ),
        "small": ParagraphStyle(
            "small",
            parent=base["BodyText"],
            fontName="MSYH",
            fontSize=8.6,
            leading=12,
            textColor=colors.HexColor("#64748B"),
            wordWrap="CJK",
        ),
        "caption": ParagraphStyle(
            "caption",
            parent=base["BodyText"],
            fontName="MSYH-Bold",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#334155"),
            alignment=TA_CENTER,
            wordWrap="CJK",
            spaceBefore=4,
            spaceAfter=8,
        ),
        "cell": ParagraphStyle(
            "cell",
            parent=base["BodyText"],
            fontName="MSYH",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#1F2937"),
            wordWrap="CJK",
        ),
        "cell_head": ParagraphStyle(
            "cell_head",
            parent=base["BodyText"],
            fontName="MSYH-Bold",
            fontSize=9.2,
            leading=13,
            textColor=colors.HexColor("#0B2545"),
            wordWrap="CJK",
        ),
    }


def p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(text.replace("\n", "<br/>"), style)


def bullet_table(items: list[str], styles: dict[str, ParagraphStyle], width: float) -> Table:
    rows = []
    for item in items:
        rows.append([p("•", styles["body"]), p(item, styles["body"])])
    tbl = Table(rows, colWidths=[0.35 * cm, width - 0.35 * cm], hAlign="LEFT")
    tbl.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )
    return tbl


def front_steps_table(styles: dict[str, ParagraphStyle], content_width: float) -> Table:
    rows = []
    for i, row in enumerate(FRONT_STEP_ROWS):
        style = styles["cell_head"] if i == 0 else styles["cell"]
        rows.append([p(cell, style) for cell in row])
    tbl = Table(rows, colWidths=[4.1 * cm, 12.0 * cm, content_width - 16.1 * cm], repeatRows=1)
    tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E0F2FE")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
                ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
            ]
        )
    )
    return tbl


def rule_table(styles: dict[str, ParagraphStyle], content_width: float) -> Table:
    rows = []
    for i, row in enumerate(RULE_ROWS):
        style = styles["cell_head"] if i == 0 else styles["cell"]
        rows.append([p(cell, style) for cell in row])
    tbl = Table(rows, colWidths=[3.1 * cm, 10.2 * cm, content_width - 13.3 * cm], repeatRows=1)
    tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
                ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
            ]
        )
    )
    return tbl


def image_flowable(path: Path, max_width: float, max_height: float | None = None) -> RLImage:
    with Image.open(path) as im:
        w, h = im.size
    ratio = h / w
    draw_w = max_width
    draw_h = draw_w * ratio
    if max_height and draw_h > max_height:
        draw_h = max_height
        draw_w = draw_h / ratio
    img = RLImage(str(path), width=draw_w, height=draw_h)
    img.hAlign = "CENTER"
    return img


def on_page(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("MSYH", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    width, height = landscape(A4)
    canvas.drawString(doc.leftMargin, height - 0.55 * cm, "附加项新定价配置说明 | 优化版")
    canvas.drawRightString(width - doc.rightMargin, 0.45 * cm, f"第 {doc.page} 页")
    canvas.setStrokeColor(colors.HexColor("#CBD5E1"))
    canvas.setLineWidth(0.4)
    canvas.line(doc.leftMargin, height - 0.72 * cm, width - doc.rightMargin, height - 0.72 * cm)
    canvas.restoreState()


def build_pdf() -> None:
    register_pdf_fonts()
    styles = pdf_styles()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=page_size,
        rightMargin=1.25 * cm,
        leftMargin=1.25 * cm,
        topMargin=1.05 * cm,
        bottomMargin=1.0 * cm,
        title=TITLE,
        author="Codex",
    )
    content_width = page_size[0] - doc.leftMargin - doc.rightMargin
    story = []

    meta = Table(
        [[p(f"<b>{TITLE}</b>", styles["title"])], [p(SUBTITLE, styles["subtitle"])]],
        colWidths=[content_width],
        hAlign="LEFT",
    )
    meta.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#CBD5E1")),
                ("LEFTPADDING", (0, 0), (-1, -1), 14),
                ("RIGHTPADDING", (0, 0), (-1, -1), 14),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.extend(
        [
            meta,
            Spacer(1, 10),
            p("一、前置配置：会员等级与会员定价", styles["h1"]),
            image_flowable(FRONT_FLOWCHART_PATH, content_width, 8.2 * cm),
            Spacer(1, 8),
            front_steps_table(styles, content_width),
            Spacer(1, 8),
            p("二、变更概述", styles["h1"]),
        ]
    )
    story.append(bullet_table(SUMMARY, styles, content_width))
    story.extend([Spacer(1, 8), p("三、核心定价规则", styles["h1"]), rule_table(styles, content_width)])
    story.append(PageBreak())

    story.extend(
        [
            p("四、配置流程图", styles["h1"]),
            image_flowable(FLOWCHART_PATH, content_width, 13.6 * cm),
            Spacer(1, 8),
            p("五、操作步骤", styles["h1"]),
        ]
    )
    for title, body in STEPS:
        story.append(p(f"<b>{title}</b>：{body}", styles["body"]))
    story.append(PageBreak())

    single_screens = SCREENSHOTS[:4]
    for idx, (caption, img_path) in enumerate(single_screens):
        story.extend(
            [
                p("六、关键界面示意" if idx == 0 else "关键界面示意（续）", styles["h1"]),
                p(caption, styles["caption"]),
                image_flowable(img_path, content_width, 14.2 * cm),
                PageBreak(),
            ]
        )

    story.extend(
        [
            p("关键界面示意（续）", styles["h1"]),
            p(SCREENSHOTS[4][0], styles["caption"]),
            image_flowable(SCREENSHOTS[4][1], content_width, 9.2 * cm),
            Spacer(1, 8),
            p(SCREENSHOTS[5][0], styles["caption"]),
            image_flowable(SCREENSHOTS[5][1], content_width, 3.8 * cm),
            PageBreak(),
            p("关键界面示意（续）", styles["h1"]),
            p(SCREENSHOTS[6][0], styles["caption"]),
            image_flowable(SCREENSHOTS[6][1], content_width, 14.2 * cm),
            PageBreak(),
            p("关键界面示意（续）", styles["h1"]),
            p(SCREENSHOTS[7][0], styles["caption"]),
            image_flowable(SCREENSHOTS[7][1], content_width, 14.2 * cm),
        ]
    )

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


def set_docx_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_docx_paragraph(paragraph, before=0, after=6, line=1.15):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def docx_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    style_docx_paragraph(p, before=10 if level == 1 else 6, after=5)
    run = p.add_run(text)
    set_docx_font(run, size=15 if level == 1 else 12, bold=True, color=RGBColor(29, 78, 216) if level == 1 else RGBColor(15, 23, 42))


def docx_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    style_docx_paragraph(p)
    run = p.add_run(text)
    set_docx_font(run, size=10.5)


def docx_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_docx_paragraph(p, after=4)
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        run = p.add_run(item)
        set_docx_font(run, size=10.5)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    header = section.header.paragraphs[0]
    header.text = "附加项新定价配置说明 | 优化版"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_docx_font(run, size=8.5, color=RGBColor(100, 116, 139))

    p_title = doc.add_paragraph()
    style_docx_paragraph(p_title, after=2)
    title_run = p_title.add_run(TITLE)
    set_docx_font(title_run, size=23, bold=True, color=RGBColor(11, 37, 69))

    p_sub = doc.add_paragraph()
    style_docx_paragraph(p_sub, after=10)
    sub_run = p_sub.add_run(f"{SUBTITLE} | {DOC_DATE}")
    set_docx_font(sub_run, size=11.5, color=RGBColor(71, 85, 105))

    docx_heading(doc, "一、前置配置：会员等级与会员定价")
    doc.add_picture(str(FRONT_FLOWCHART_PATH), width=Inches(9.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(FRONT_STEP_ROWS[0]):
        hdr[i].text = h
        shade_cell(hdr[i], "E0F2FE")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in hdr[i].paragraphs[0].runs:
            set_docx_font(run, size=9.5, bold=True, color=RGBColor(11, 37, 69))
    for row in FRONT_STEP_ROWS[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[i].paragraphs[0].runs:
                set_docx_font(run, size=9)

    docx_heading(doc, "二、变更概述")
    docx_bullets(doc, SUMMARY)

    docx_heading(doc, "三、核心定价规则")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(RULE_ROWS[0]):
        hdr[i].text = h
        shade_cell(hdr[i], "E8EEF5")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in hdr[i].paragraphs[0].runs:
            set_docx_font(run, size=9.5, bold=True, color=RGBColor(11, 37, 69))
    for row in RULE_ROWS[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[i].paragraphs[0].runs:
                set_docx_font(run, size=9)

    doc.add_page_break()
    docx_heading(doc, "四、配置流程图")
    doc.add_picture(str(FLOWCHART_PATH), width=Inches(9.9))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    docx_heading(doc, "五、操作步骤")
    for title, body in STEPS:
        p = doc.add_paragraph()
        style_docx_paragraph(p, after=5)
        r1 = p.add_run(f"{title}：")
        set_docx_font(r1, size=10.5, bold=True)
        r2 = p.add_run(body)
        set_docx_font(r2, size=10.5)

    doc.add_page_break()
    docx_heading(doc, "六、关键界面示意")
    for idx, (caption, img_path) in enumerate(SCREENSHOTS):
        p = doc.add_paragraph()
        style_docx_paragraph(p, before=4, after=4)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption)
        set_docx_font(r, size=9.5, bold=True, color=RGBColor(51, 65, 85))
        doc.add_picture(str(img_path), width=Inches(9.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if idx in {1, 3, 5, 7} and idx != len(SCREENSHOTS) - 1:
            doc.add_page_break()

    doc.save(DOCX_PATH)


def render_pdf_pages() -> None:
    for old in RENDER_DIR.glob("page-*.png"):
        old.unlink()
    pdf = fitz.open(PDF_PATH)
    for i, page in enumerate(pdf, start=1):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        pix.save(RENDER_DIR / f"page-{i:02d}.png")
    pdf.close()


def main() -> None:
    ensure_dirs()
    generate_front_flowchart()
    generate_flowchart()
    build_pdf()
    build_docx()
    render_pdf_pages()
    print(f"PDF: {PDF_PATH}")
    print(f"DOCX: {DOCX_PATH}")
    print(f"FLOWCHART: {FLOWCHART_PATH}")
    print(f"RENDER_DIR: {RENDER_DIR}")


if __name__ == "__main__":
    main()
