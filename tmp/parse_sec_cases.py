import openpyxl
from collections import defaultdict

path = r"D:\test_workspace\安全优化\测试文档\special_安全优化测试用例_20260624_084827.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb[wb.sheetnames[0]]
prefix_map = {
    "TC-SEC-FIN": "资金与支付安全",
    "TC-SEC-AUTH": "鉴权与访问控制",
    "TC-SEC-IDOR": "资源归属与 IDOR",
    "TC-SEC-IM": "IM 与消息安全",
    "TC-SEC-XSS": "XSS 与内容注入",
    "TC-SEC-SECRET": "密钥与敏感数据",
    "TC-SEC-EXT": "Chrome 扩展与本地组件",
    "TC-SEC-AUD": "审计与发布门禁",
}
groups = defaultdict(lambda: {"pass": 0, "block": 0, "total": 0, "blocked": []})
for r in range(2, ws.max_row + 1):
    title = str(ws.cell(r, 3).value or "")
    status = str(ws.cell(r, 5).value or "")
    domain = "其他"
    for prefix, name in prefix_map.items():
        if prefix in title:
            domain = name
            break
    groups[domain]["total"] += 1
    if status == "通过":
        groups[domain]["pass"] += 1
    else:
        groups[domain]["block"] += 1
        groups[domain]["blocked"].append(title.split("]")[0] + "]")

for domain, data in groups.items():
    print(domain, f"{data['pass']}/{data['total']}", data["blocked"])
